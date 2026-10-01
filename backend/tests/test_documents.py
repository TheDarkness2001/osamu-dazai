"""Document output: DOCX template engine, HTML, PDF, figure conversion, pre-export checks."""

import asyncio
import zipfile
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

import docx
import pytest

from osamu_dazai.documents.checks import check_student_book
from osamu_dazai.documents.docx_render import assessment_docx, student_book_docx, teacher_guide_docx
from osamu_dazai.documents.figures import Figure, svg_to_png
from osamu_dazai.documents.html import assessment_html, student_book_html, teacher_guide_html
from osamu_dazai.documents.inline import Run, parse_inline, to_html, unbalanced
from osamu_dazai.documents.pdf import find_chromium, html_to_pdf
from osamu_dazai.domain.common import Lang
from osamu_dazai.domain.knowledge import KnowledgeBase
from osamu_dazai.pipeline.book import STAGE as BOOK_STAGE
from osamu_dazai.pipeline.book import BlockDraft, BookWriter, plan_chapters
from osamu_dazai.pipeline.curriculum import assemble_course, build_graph
from osamu_dazai.pipeline.teacher import STAGE as TEACHER_STAGE
from osamu_dazai.pipeline.teacher import TeacherGuideWriter
from osamu_dazai.providers import ProviderRegistry
from osamu_dazai.providers.image import MockImageProvider
from osamu_dazai.providers.llm import MockProvider
from osamu_dazai.questions.generator import assemble, blueprint
from osamu_dazai.questions.shuffle import make_versions
from osamu_dazai.storage.assets import AssetStore
from osamu_dazai.visuals.builder import STAGE as VIS_STAGE
from osamu_dazai.visuals.builder import VisualBuilder
from tests.fixtures.assessment_sample import assessment_draft
from tests.fixtures.book_sample import chapter_draft
from tests.fixtures.course_sample import curriculum_draft, graph_draft
from tests.fixtures.teacher_sample import teacher_draft
from tests.test_course_design import brief
from tests.test_visuals import plan_for

run = asyncio.run
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def reg_with(stage, replies, images=False):
    r = ProviderRegistry()
    r.register_llm(MockProvider().queue(stage, *[x if isinstance(x, str) else x.model_dump_json() for x in replies]))
    if images:
        r.register_image(MockImageProvider())
    return r


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("course")
    graph, ids = build_graph(graph_draft(12), "c")
    course = assemble_course(curriculum_draft(12, 3, 6), brief(), "prj_1", graph, ids)
    plans = plan_chapters(course, graph)
    drafts = [chapter_draft(p, graph) for p in plans]
    drafts[1].sections[3].blocks.append(BlockDraft(
        kind="visual", visual_kind="illustration", title="Kompyuter oldidagi o‘quvchi",
        visual_purpose="mavzuga qiziqish uyg‘otish",
        visual_description="Bir o‘quvchi maktab kompyuter xonasida dastur yozmoqda, ekranda rangli bloklar"))
    book = run(BookWriter(reg_with(BOOK_STAGE, drafts)).write(
        course, graph, KnowledgeBase(project_id="p"), words_per_chapter=40)).book
    res = run(VisualBuilder(reg_with(VIS_STAGE, plan_for(book), images=True), AssetStore(tmp)).build(
        book, course, graph, KnowledgeBase(project_id="p")))
    figures = {vid: Figure(vid, tmp / rel) for vid, rel in res.figures.items()}
    guide = run(TeacherGuideWriter(reg_with(TEACHER_STAGE, [teacher_draft(p) for p in plans])).write(
        course, graph, res.book, KnowledgeBase(project_id="p"))).guide
    bp = blueprint(course, graph, course.assessment_plan[0], mix={"multiple_choice": 3, "true_false": 1,
                                                                  "short_answer": 1})
    quiz = assemble(assessment_draft(bp), bp, course)
    return course, res.book, figures, guide, quiz


# ---- inline + figures --------------------------------------------------------------
def test_inline_markup():
    assert parse_inline("a **b** *c* `d`") == [Run("a "), Run("b", bold=True), Run(" "), Run("c", italic=True),
                                                Run(" "), Run("d", code=True)]
    assert to_html("x < y **and** `<b>`") == "x &lt; y <strong>and</strong> <code>&lt;b&gt;</code>"
    assert parse_inline("2 * 3 * 4") == [Run("2 * 3 * 4")]  # arithmetic is not italics
    assert unbalanced("a **b") and not unbalanced("a **b**")


def test_svg_to_png(built):
    _, _, figures, _, _ = built
    png = svg_to_png(next(iter(figures.values())).svg_text())
    assert png.startswith(b"\x89PNG")


# ---- DOCX -----------------------------------------------------------------------------
def doc_xml(data: bytes) -> str:
    with zipfile.ZipFile(BytesIO(data)) as z:
        return z.read("word/document.xml").decode("utf-8")


def test_student_book_docx(built):
    course, book, figures, _, _ = built
    data = student_book_docx(book, figures=figures)
    d = docx.Document(BytesIO(data))
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Mundarija" in headings and f"Bob 1. {book.chapters[0].title}" in headings
    assert "Lug‘at" in headings and "Javoblar" in headings
    assert len(d.inline_shapes) == len(figures) == 4
    xml = doc_xml(data)
    assert 'w:instr="TOC' in xml and "Osamu Dazai Code" in {s.name for s in d.styles}
    with zipfile.ZipFile(BytesIO(data)) as z:
        footer = next(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
        header = next(z.read(n).decode() for n in z.namelist() if n.startswith("word/header"))
    assert 'w:instr="PAGE"' in footer and course.title in header
    sec = d.sections[0]
    assert round(sec.page_width.mm) == 210 and round(sec.page_height.mm) == 297
    assert "uz-Latn-UZ" in d.styles.element.xml
    body_text = "\n".join(p.text for p in d.paragraphs)
    assert "Rasm 1.1: Dastur qadamlari" in body_text


def test_student_edition_has_no_answers(built):
    _, book, figures, _, _ = built
    with_key = "\n".join(p.text for p in docx.Document(BytesIO(student_book_docx(book, figures=figures))).paragraphs)
    without = "\n".join(p.text for p in docx.Document(BytesIO(student_book_docx(book, answer_key=False))).paragraphs)
    assert "Ekranga matn chiqaradi." in with_key and "Ekranga matn chiqaradi." not in without


def test_teacher_guide_docx(built):
    course, book, _, guide, _ = built
    d = docx.Document(BytesIO(teacher_guide_docx(guide, course, book)))
    texts = "\n".join(p.text for p in d.paragraphs)
    assert "O‘qituvchi uchun qo‘llanma" in texts and "Dars 1. Lesson 1 (45 daq)" in texts
    timing = [t for t in d.tables if t.rows[0].cells[1].text == "Faoliyat"]
    assert len(timing) == 18 and len(timing[0].rows) == 4  # header + 3 timed steps
    keys = [t for t in d.tables if t.rows[0].cells[0].text == "#"]
    assert len(keys) == 3


def test_assessment_docx_versions(built):
    _, _, _, _, quiz = built
    v = make_versions(quiz, 2)[1]
    student = "\n".join(p.text for p in docx.Document(BytesIO(assessment_docx(v, Lang.UZ))).paragraphs)
    key = "\n".join(p.text for p in docx.Document(BytesIO(assessment_docx(v, Lang.UZ, answer_key=True))).paragraphs)
    assert "Variant B" in student and "Ism: ____" in student and "☐ To‘g‘ri" in student
    assert "Birinchi variant to‘g‘ri" not in student and "Birinchi variant to‘g‘ri" in key
    assert key.startswith("Javoblar kaliti:")


# ---- HTML ---------------------------------------------------------------------------------
class _Checker(HTMLParser):
    VOID = {"meta", "br", "img", "rect", "line", "circle", "path", "polygon", "polyline"}

    def __init__(self) -> None:
        super().__init__()
        self.stack: list[str] = []
        self.errors: list[str] = []
        self.scripts = 0

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag == "script":
            self.scripts += 1
        if tag not in self.VOID:
            self.stack.append(tag)

    def handle_startendtag(self, tag, attrs):  # noqa: ANN001
        pass

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag in self.VOID:
            return
        if not self.stack or self.stack[-1] != tag:
            self.errors.append(f"unexpected </{tag}> (open: {self.stack[-3:]})")
        else:
            self.stack.pop()


def well_formed(html: str) -> _Checker:
    c = _Checker()
    c.feed(html)
    return c


def test_student_book_html(built):
    _, book, figures, _, _ = built
    evil = book.model_copy(deep=True)
    evil.chapters[0].sections[1].blocks[0].text = "Safe? <script>alert(1)</script> **yes**"
    html = student_book_html(evil, figures=figures)
    chk = well_formed(html)
    assert chk.errors == [] and chk.stack == [] and chk.scripts == 0
    assert '<html lang="uz">' in html and "&lt;script&gt;" in html and "<strong>yes</strong>" in html
    assert html.count("<figure><svg") == 3 and html.count("data:image/png;base64") == 1
    assert "@media print" in html and 'href="#ch2"' in html


def test_teacher_and_assessment_html(built):
    course, book, _, guide, quiz = built
    for html in (teacher_guide_html(guide, course, book), assessment_html(quiz, Lang.UZ),
                 assessment_html(quiz, Lang.UZ, answer_key=True)):
        chk = well_formed(html)
        assert chk.errors == [] and chk.stack == []
    assert "Test savollari" in assessment_html(quiz, Lang.UZ)


# ---- PDF ------------------------------------------------------------------------------------
@pytest.mark.skipif(find_chromium() is None, reason="no Chromium-based browser available")
def test_html_to_pdf(built, tmp_path):
    _, book, figures, _, _ = built
    pdf = html_to_pdf(student_book_html(book, figures=figures), tmp_path / "book.pdf")
    data = Path(pdf).read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 5000


# ---- checks ---------------------------------------------------------------------------------
def test_document_checks(built, tmp_path):
    _, book, figures, _, _ = built
    assert check_student_book(book, figures).passed
    assert check_student_book(book, figures).issues == []
    bad = book.model_copy(deep=True)
    bad.chapters[1].number = 5
    bad.chapters[0].visual_specs = {}
    bad.chapters[0].exercises[0].questions[1].number = 7
    bad.chapters[2].sections[1].blocks[0].text = "Bu **qalin matn"
    codes = {i.code for i in check_student_book(bad, figures).issues}
    assert {"document.chapter_numbering", "document.broken_figure_ref", "document.exercise_numbering",
            "document.broken_formatting"} <= codes
    missing = {k: Figure(k, tmp_path / "nope.svg") for k in figures}
    assert "document.missing_image" in {i.code for i in check_student_book(book, missing).errors}
    unrendered = check_student_book(book, {}, require_figures=True)
    assert not unrendered.passed
