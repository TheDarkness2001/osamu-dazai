"""Helpers to build realistic *broken* Word documents for IELTS tests."""

from __future__ import annotations

from io import BytesIO

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from osamu_dazai.ielts import lines_to_text, render_test
from tests.fixtures.ielts_sample import sample_test


def sample_text() -> str:
    return lines_to_text(render_test(sample_test()))


def sample_lines() -> list[str]:
    return sample_text().split("\n")


def _add_roman_numbering(d: docx.document.Document) -> str:
    """Register a lowerRoman list ("i.", "ii." …) and return its numId."""
    numbering = d.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/>'
        '<w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'
    )
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>')
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(num)
    return "90"


def build_docx(
    lines: list[str],
    *,
    roman_auto: set[int] | None = None,  # line indexes rendered with Word auto-numbering
    soft_break_with_next: set[int] | None = None,  # join line i and i+1 with Shift+Enter
    table_lines: set[int] | None = None,  # put these lines in a one-column table
) -> bytes:
    roman_auto = roman_auto or set()
    soft = soft_break_with_next or set()
    table_lines = table_lines or set()
    d = docx.Document()
    num_id = _add_roman_numbering(d) if roman_auto else None
    table = None
    i = 0
    while i < len(lines):
        text = lines[i]
        if i in table_lines:
            if table is None:
                table = d.add_table(rows=0, cols=1)
            table.add_row().cells[0].text = text
            i += 1
            continue
        table = None
        p = d.add_paragraph()
        run = p.add_run(text)
        if i in soft and i + 1 < len(lines):
            run.add_break()
            p.add_run(lines[i + 1])
            i += 1
        if i in roman_auto:
            ppr = p._p.get_or_add_pPr()
            ppr.append(parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
            ))
        i += 1
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()
