"""DocumentTemplateEngine (spec §35): structured content → formatted DOCX.

Page size, margins, fonts (Latin + Cyrillic), heading colours, line spacing,
running header, page-number footer, table of contents, headings, inline
bold/italic/code, lists, code blocks, callouts, tables, figures with captions,
page breaks. IELTS parser documents do NOT use this engine — they use the
parser-safe writer in ``osamu_dazai.ielts.docx_io``.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

from osamu_dazai.documents.figures import Figure
from osamu_dazai.documents.inline import parse_inline
from osamu_dazai.documents.template import PAGE_SIZES_MM, DocTemplate

CALLOUT_FILL = {"note": "E8F1FB", "tip": "EAF7EE", "warning": "FFF4E0", "common_mistake": "FDECEC",
                "definition": "EEF0FA", "example": "F4F4F4"}


def _set_fonts(style, font: str) -> None:  # noqa: ANN001
    style.font.name = font
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)


def _shade(cell_or_par_pr, fill: str) -> None:  # noqa: ANN001
    cell_or_par_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def _field(paragraph, instr: str, placeholder: str = "") -> None:  # noqa: ANN001
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr={quoteattr(instr)}><w:r><w:t xml:space="preserve">'
                    f'{escape(placeholder)}</w:t></w:r></w:fldSimple>')
    paragraph._p.append(fld)


class DocumentTemplateEngine:
    def __init__(self, template: DocTemplate, *, title: str = "", lang: str = "en") -> None:
        self.t = template
        self.doc = docx.Document()
        self.title = title
        self.lang = lang
        self._setup()

    # ---- setup -----------------------------------------------------------
    def _setup(self) -> None:
        t, d = self.t, self.doc
        w, h = PAGE_SIZES_MM[t.page_size]
        sec = d.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Mm(w), Mm(h)
        sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = (Mm(x) for x in t.margins_mm)
        normal = d.styles["Normal"]
        _set_fonts(normal, t.body_font)
        normal.font.size = Pt(t.base_pt)
        normal.paragraph_format.line_spacing = t.line_spacing
        normal.paragraph_format.space_after = Pt(6)
        # language tag for spell-checking / hyphenation
        lang_el = OxmlElement("w:lang")
        lang_el.set(qn("w:val"), {"uz": "uz-Latn-UZ", "uz-Cyrl": "uz-Cyrl-UZ", "ru": "ru-RU"}.get(self.lang, "en-GB"))
        normal.element.get_or_add_rPr().append(lang_el)
        for level, size in ((1, t.base_pt + 9), (2, t.base_pt + 5), (3, t.base_pt + 2.5), (4, t.base_pt + 1)):
            st = d.styles[f"Heading {level}"]
            _set_fonts(st, t.heading_font)
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor.from_string(t.heading_color)
            st.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        code = d.styles.add_style("Osamu Dazai Code", 1)
        _set_fonts(code, t.code_font)
        code.font.size = Pt(t.base_pt - 1.5)
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.line_spacing = 1.0
        caption = d.styles["Caption"]
        caption.font.size = Pt(t.base_pt - 1)
        caption.font.italic = True
        if t.header_text:
            hp = sec.header.paragraphs[0]
            hp.text = t.header_text.replace("{title}", self.title)
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in hp.runs:
                r.font.size, r.font.color.rgb = Pt(t.base_pt - 2), RGBColor(0x7B, 0x87, 0x94)
        if t.page_numbers:
            fp = sec.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _field(fp, "PAGE", "1")
        settings = d.settings.element
        settings.append(parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'))  # refresh TOC/fields on open

    # ---- building blocks -------------------------------------------------
    def runs(self, paragraph, text: str, *, bold: bool = False) -> None:  # noqa: ANN001
        for r in parse_inline(text):
            run = paragraph.add_run(r.text)
            run.bold = r.bold or bold or None
            run.italic = r.italic or None
            if r.code:
                run.font.name = self.t.code_font
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), self.t.code_font)

    def title_block(self, title: str, subtitle: str = "") -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.bold, r.font.size = True, Pt(self.t.base_pt + 15)
        r.font.color.rgb = RGBColor.from_string(self.t.heading_color)
        if subtitle:
            s = self.doc.add_paragraph()
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.add_run(subtitle).font.size = Pt(self.t.base_pt + 3)

    def toc(self, heading: str) -> None:
        if not self.t.table_of_contents:
            return
        self.heading(heading, 1)
        _field(self.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "…")
        self.page_break()

    def heading(self, text: str, level: int) -> None:
        self.doc.add_heading(text, level=min(level, 4))

    def paragraph(self, text: str, *, bold: bool = False, italic: bool = False, indent_cm: float = 0) -> None:
        p = self.doc.add_paragraph()
        self.runs(p, text, bold=bold)
        if italic:
            for r in p.runs:
                r.italic = True
        if indent_cm:
            p.paragraph_format.left_indent = Cm(indent_cm)

    def bullets(self, items: list[str], *, ordered: bool = False) -> None:
        for i, item in enumerate(items, start=1):
            if ordered:  # literal numbers: Word's numbered lists continue across lists unpredictably
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.6)
                self.runs(p, f"{i}.\t{item}")
            else:
                p = self.doc.add_paragraph(style="List Bullet")
                self.runs(p, item)

    def code(self, code: str, *, output: str | None = None) -> None:
        for block, fill in ((code, "F3F4F6"), (output, "FFFFFF")):
            if block is None:
                continue
            table = self.doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            cell = table.rows[0].cells[0]
            _shade(cell._tc.get_or_add_tcPr(), fill)
            cell.paragraphs[0].style = self.doc.styles["Osamu Dazai Code"]
            lines = block.rstrip("\n").split("\n")
            cell.paragraphs[0].add_run(lines[0])
            for ln in lines[1:]:
                cell.add_paragraph(ln, style="Osamu Dazai Code")
            self.doc.add_paragraph()

    def callout(self, style: str, title: str, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        _shade(cell._tc.get_or_add_tcPr(), CALLOUT_FILL.get(style, "F4F4F4"))
        p = cell.paragraphs[0]
        if title:
            p.add_run(title).bold = True
            p = cell.add_paragraph()
        self.runs(p, text)
        self.doc.add_paragraph()

    def table(self, header: list[str], rows: list[list[str]], *, caption: str = "",
              widths_cm: list[float] | None = None) -> None:
        if caption:
            self.doc.add_paragraph(caption, style="Caption")
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = table.rows[0].cells[i]
            _shade(c._tc.get_or_add_tcPr(), "E3F2FD")
            self.runs(c.paragraphs[0], h, bold=True)
        # repeat header row on each page
        table.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for row in rows:
            cells = table.add_row().cells
            for i, v in enumerate(row):
                self.runs(cells[i].paragraphs[0], v)
        if widths_cm:
            for row in table.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()

    def figure(self, fig: Figure, caption: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(BytesIO(fig.png_bytes()), width=Cm(fig.width_cm(self.t.text_width_cm)))
        cap = self.doc.add_paragraph(caption, style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def figure_placeholder(self, caption: str) -> None:
        self.paragraph(f"[{caption}]", italic=True)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, dest: str | Path | BytesIO) -> None:
        self.doc.save(dest)

    def to_bytes(self) -> bytes:
        buf = BytesIO()
        self.doc.save(buf)
        return buf.getvalue()
