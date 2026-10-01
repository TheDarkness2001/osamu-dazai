"""Parser-safe DOCX writing and faithful DOCX reading for IELTS documents.

Writer: one ``w:p`` per line, plain runs, no list styles, no numbering, no
tables, no fields (spec §35 "parser compatibility over decoration").

Reader: rebuilds the line IR from ``w:p`` elements in body order, including
paragraphs inside tables, and recovers what the page *shows* but the text
does not contain — Word auto-numbering labels — plus soft line breaks.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

from osamu_dazai.ielts.lines import Line

Source = str | Path | bytes | BinaryIO


@dataclass(frozen=True)
class DocxStyle:
    font: str = "Times New Roman"
    size_pt: float = 12
    space_after_pt: float = 4


def write_docx(lines: list[Line], dest: str | Path | BinaryIO, style: DocxStyle = DocxStyle()) -> None:
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = style.font
    normal.font.size = Pt(style.size_pt)
    normal.paragraph_format.space_after = Pt(style.space_after_pt)
    for ln in lines:
        if "\n" in ln.text:
            raise ValueError(f"line contains a newline; split it first: {ln.text[:60]!r}")
        p = d.add_paragraph()
        if ln.text:
            run = p.add_run(ln.text)
            run.bold = ln.bold or None
    d.save(dest)


def docx_bytes(lines: list[Line], style: DocxStyle = DocxStyle()) -> bytes:
    buf = BytesIO()
    write_docx(lines, buf, style)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------
def _roman(n: int) -> str:
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
            (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out, n = out + s, n - v
    return out


def _letters(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(ord("a") + r) + s
    return s


def _fmt(n: int, fmt: str) -> str:
    return {
        "lowerRoman": _roman(n),
        "upperRoman": _roman(n).upper(),
        "lowerLetter": _letters(n),
        "upperLetter": _letters(n).upper(),
        "decimalZero": f"{n:02d}",
        "bullet": "",
        "none": "",
    }.get(fmt, str(n))


class _Numbering:
    """Computes the labels Word would display for auto-numbered paragraphs."""

    def __init__(self, document: docx.document.Document) -> None:
        self.levels: dict[str, dict[int, tuple[str, str, int]]] = {}  # abstractId → ilvl → (fmt, text, start)
        self.nums: dict[str, tuple[str, dict[int, int]]] = {}  # numId → (abstractId, start overrides)
        self.style_num: dict[str, tuple[str, int]] = {}
        self.counters: dict[str, list[int]] = {}
        try:
            root = document.part.numbering_part.element
        except (KeyError, NotImplementedError, AttributeError):
            root = None
        if root is not None:
            for ab in root.findall(qn("w:abstractNum")):
                lv = {}
                for lvl in ab.findall(qn("w:lvl")):
                    fmt = lvl.find(qn("w:numFmt"))
                    txt = lvl.find(qn("w:lvlText"))
                    start = lvl.find(qn("w:start"))
                    lv[int(lvl.get(qn("w:ilvl")))] = (
                        fmt.get(qn("w:val")) if fmt is not None else "decimal",
                        txt.get(qn("w:val")) if txt is not None else "%1.",
                        int(start.get(qn("w:val"))) if start is not None else 1,
                    )
                self.levels[ab.get(qn("w:abstractNumId"))] = lv
            for num in root.findall(qn("w:num")):
                ab = num.find(qn("w:abstractNumId")).get(qn("w:val"))
                over = {}
                for o in num.findall(qn("w:lvlOverride")):
                    so = o.find(qn("w:startOverride"))
                    if so is not None:
                        over[int(o.get(qn("w:ilvl")))] = int(so.get(qn("w:val")))
                self.nums[num.get(qn("w:numId"))] = (ab, over)
        for st in document.styles.element.findall(qn("w:style")):
            ppr = st.find(qn("w:pPr"))
            npr = ppr.find(qn("w:numPr")) if ppr is not None else None
            if npr is not None and npr.find(qn("w:numId")) is not None:
                ilvl = npr.find(qn("w:ilvl"))
                self.style_num[st.get(qn("w:styleId"))] = (
                    npr.find(qn("w:numId")).get(qn("w:val")),
                    int(ilvl.get(qn("w:val"))) if ilvl is not None else 0,
                )

    def label(self, p) -> str | None:  # noqa: ANN001 - lxml element
        ppr = p.find(qn("w:pPr"))
        num_id, ilvl = None, 0
        if ppr is not None:
            npr = ppr.find(qn("w:numPr"))
            if npr is not None and npr.find(qn("w:numId")) is not None:
                num_id = npr.find(qn("w:numId")).get(qn("w:val"))
                il = npr.find(qn("w:ilvl"))
                ilvl = int(il.get(qn("w:val"))) if il is not None else 0
            elif (ps := ppr.find(qn("w:pStyle"))) is not None and ps.get(qn("w:val")) in self.style_num:
                num_id, ilvl = self.style_num[ps.get(qn("w:val"))]
        if num_id is None or num_id == "0":
            return None
        ab, over = self.nums.get(num_id, (None, {}))
        lv = self.levels.get(ab, {}) if ab is not None else {}
        cnt = self.counters.setdefault(num_id, [0] * 9)
        if cnt[ilvl] == 0:
            cnt[ilvl] = over.get(ilvl, lv.get(ilvl, ("decimal", "%1.", 1))[2])
        else:
            cnt[ilvl] += 1
        for deeper in range(ilvl + 1, 9):
            cnt[deeper] = 0
        fmt, text, _ = lv.get(ilvl, ("decimal", f"%{ilvl + 1}.", 1))
        if fmt == "bullet":
            return text or "•"
        for k in range(9, 0, -1):
            if f"%{k}" in text:
                kfmt = lv.get(k - 1, ("decimal", "", 1))[0]
                text = text.replace(f"%{k}", _fmt(max(cnt[k - 1], 1), kfmt))
        return text.strip() or "•"


def _paragraph_text(p) -> tuple[str, int]:  # noqa: ANN001
    parts: list[str] = []
    breaks = 0
    for el in p.iter():
        if el.tag == qn("w:t"):
            parts.append(el.text or "")
        elif el.tag == qn("w:tab"):
            parts.append("\t")
        elif el.tag in (qn("w:br"), qn("w:cr")):
            if el.get(qn("w:type")) in (None, "textWrapping"):
                parts.append("\n")
                breaks += 1
    return "".join(parts), breaks


@dataclass
class DocxRead:
    lines: list[Line]
    tables: int


def _open(src: Source) -> docx.document.Document:
    if isinstance(src, bytes):
        return docx.Document(BytesIO(src))
    return docx.Document(src)


def read_docx(src: Source) -> DocxRead:
    d = _open(src)
    numbering = _Numbering(d)
    lines: list[Line] = []
    tables = 0
    para_index = 0

    def add_p(p, in_table: bool) -> None:  # noqa: ANN001
        nonlocal para_index
        text, breaks = _paragraph_text(p)
        lines.append(Line(text=text, para_index=para_index, auto_label=numbering.label(p),
                          soft_breaks=breaks, in_table=in_table))
        para_index += 1

    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            add_p(child, False)
        elif child.tag == qn("w:tbl"):
            tables += 1
            for p in child.iter(qn("w:p")):
                add_p(p, True)
    return DocxRead(lines, tables)
